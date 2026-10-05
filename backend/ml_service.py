"""Bounded tabular training and inference for application-generated models."""

from __future__ import annotations

import csv
import json
import logging
import os
import platform
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from pathlib import Path
from uuid import uuid4

import joblib
import numpy as np
import pandas as pd
import scipy
import sklearn
from fastapi import HTTPException
from openpyxl import load_workbook
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import RandomForestClassifier, RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.inspection import permutation_importance
from sklearn.linear_model import LogisticRegression, Ridge
from sklearn.metrics import (
    accuracy_score,
    confusion_matrix,
    f1_score,
    mean_absolute_error,
    mean_squared_error,
    precision_score,
    r2_score,
    recall_score,
)
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from .config import get_settings
from .database import make_engine
from .models import Dataset, ModelMetadata, TrainingConfig, TrainingRun, utcnow

MAX_TRAIN_ROWS = 100_000
MAX_FEATURES = 50
MAX_CATEGORIES = 200
MAX_BATCH_ROWS = 5_000
MAX_BATCH_BYTES = 5 * 1024 * 1024
ARTIFACT_NAME = re.compile(r"[0-9a-f]{32}\.joblib\Z")
ALGORITHMS = {
    "classification": {"logistic_regression", "random_forest_classifier"},
    "regression": {"ridge_regression", "random_forest_regressor"},
}
executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="mindmesh-training")
logger = logging.getLogger(__name__)


def _json(value) -> str:
    return json.dumps(value, allow_nan=False, default=lambda item: item.item())


def _missing(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _stored_path(dataset: Dataset) -> Path:
    if not dataset.stored_name or not re.fullmatch(
        r"[0-9a-f]{32}\.(csv|xlsx)", dataset.stored_name
    ):
        raise ValueError("This preview-only dataset has no original file. Re-upload it to train.")
    path = get_settings().upload_dir / dataset.stored_name
    if not path.is_file():
        raise ValueError("The original dataset file is unavailable. Re-upload it to train.")
    return path


def load_frame(dataset: Dataset) -> pd.DataFrame:
    """Read the complete selected dataset while retaining text such as leading-zero IDs."""
    path = _stored_path(dataset)
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            rows = list(csv.reader(source, strict=True))
    else:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook[dataset.sheet_name or workbook.sheetnames[0]]
            rows = []
            for raw in sheet.iter_rows(values_only=False):
                row = []
                for cell in raw:
                    if cell.data_type == "f":
                        raise ValueError("Formula cells are not supported. Export values first.")
                    value = cell.value
                    if isinstance(value, bool):
                        value = str(value).lower()
                    elif isinstance(value, (date, datetime)):
                        value = value.isoformat()
                    row.append("" if value is None else value)
                rows.append(row)
        finally:
            workbook.close()
    rows = [row for row in rows if any(not _missing(value) for value in row)]
    if len(rows) < 2:
        raise ValueError("The dataset needs a header and at least one data row.")
    headers = [str(value).strip() for value in rows[0]]
    if any(len(row) != len(headers) for row in rows[1:]):
        raise ValueError("A data row no longer matches the dataset header.")
    return pd.DataFrame(rows[1:], columns=headers, dtype=object)


def dataset_schema(dataset: Dataset) -> dict:
    columns = json.loads(dataset.columns_json)
    output = []
    for column in columns:
        warnings = []
        name = column["name"]
        lowered = name.casefold()
        if column["missing"] == dataset.row_count:
            warnings.append("All values are missing; exclude this column.")
        if column.get("uniqueCountCapped") or column.get("uniqueCount", 0) > 30:
            warnings.append("High cardinality may indicate free text or an identifier.")
        if any(word in lowered for word in ("id", "code", "number", "uuid")):
            warnings.append("This looks identifier-like and may not generalize.")
        if column["type"] in {"date", "text", "mixed"}:
            warnings.append(
                f"{column['type'].title()} features are not supported in this demo; exclude or convert it."
            )
        output.append({**column, "warnings": warnings})
    return {
        "datasetId": dataset.id,
        "name": dataset.name,
        "rowCount": dataset.row_count,
        "columns": output,
        "supportedFeatureTypes": ["number", "category", "boolean"],
    }


def validate_config(dataset: Dataset, config: dict) -> None:
    task = config["task"]
    algorithms = config["algorithms"]
    if task not in ALGORITHMS:
        raise HTTPException(400, "Choose classification or regression.")
    if not algorithms or len(algorithms) > 2 or set(algorithms) - ALGORITHMS[task]:
        raise HTTPException(400, "Choose one or both algorithms available for this task.")
    names = {item["name"]: item for item in json.loads(dataset.columns_json)}
    target = config["target"]
    features = config["features"]
    if target not in names:
        raise HTTPException(400, "Choose a target column from this dataset.")
    if not features:
        raise HTTPException(400, "Choose at least one input feature.")
    if len(features) > MAX_FEATURES:
        raise HTTPException(400, f"Choose no more than {MAX_FEATURES} input features.")
    if (
        len(features) != len(set(features))
        or target in features
        or any(item not in names for item in features)
    ):
        raise HTTPException(
            400, "Features must be unique dataset columns and cannot include the target."
        )
    unsupported = [
        name for name in features if names[name]["type"] not in {"number", "category", "boolean"}
    ]
    if unsupported:
        raise HTTPException(
            400,
            f"Exclude or convert unsupported feature columns: {', '.join(unsupported)}.",
        )
    if task == "regression" and names[target]["type"] != "number":
        raise HTTPException(400, "Regression needs a numeric target column.")
    if not 0.1 <= config["testSize"] <= 0.4:
        raise HTTPException(400, "The held-out test split must be between 10% and 40%.")


def _pipeline(task: str, algorithm: str, numeric: list[str], categorical: list[str], seed: int):
    numeric_steps = [("impute", SimpleImputer(strategy="median"))]
    if algorithm in {"logistic_regression", "ridge_regression"}:
        numeric_steps.append(("scale", StandardScaler()))
    preprocessing = ColumnTransformer(
        [
            ("numeric", Pipeline(numeric_steps), numeric),
            (
                "categorical",
                Pipeline(
                    [
                        ("impute", SimpleImputer(strategy="most_frequent")),
                        ("encode", OneHotEncoder(handle_unknown="ignore")),
                    ]
                ),
                categorical,
            ),
        ],
        remainder="drop",
    )
    if algorithm == "logistic_regression":
        estimator = LogisticRegression(max_iter=500, random_state=seed)
    elif algorithm == "random_forest_classifier":
        estimator = RandomForestClassifier(n_estimators=100, random_state=seed, n_jobs=1)
    elif algorithm == "ridge_regression":
        estimator = Ridge(alpha=1.0)
    else:
        estimator = RandomForestRegressor(n_estimators=100, random_state=seed, n_jobs=1)
    return Pipeline([("prepare", preprocessing), ("model", estimator)])


def _distribution(values) -> dict[str, int]:
    counts = pd.Series(values).value_counts()
    return {str(key): int(value) for key, value in counts.items()}


def _importance(pipeline, features, X_test, y_test, task, seed):
    sample = X_test.iloc[:1000]
    labels = y_test.iloc[:1000]
    scoring = "f1_macro" if task == "classification" else "neg_mean_absolute_error"
    measured = permutation_importance(
        pipeline, sample, labels, n_repeats=2, random_state=seed, scoring=scoring, n_jobs=1
    )
    return [
        {"feature": name, "importance": float(value)}
        for name, value in sorted(
            zip(features, measured.importances_mean, strict=True),
            key=lambda item: item[1],
            reverse=True,
        )
    ]


def train_models(
    dataset: Dataset, config: dict
) -> tuple[list[dict], list[tuple[str, object, dict]]]:
    frame = load_frame(dataset)
    target = config["target"]
    features = config["features"]
    task = config["task"]
    seed = config["randomSeed"]
    target_missing = frame[target].map(_missing)
    omitted = int(target_missing.sum())
    frame = frame.loc[~target_missing, [*features, target]].copy()
    if len(frame) > MAX_TRAIN_ROWS:
        raise ValueError(f"Training is limited to {MAX_TRAIN_ROWS:,} usable rows.")
    if task == "regression":
        numeric_target = pd.to_numeric(frame[target], errors="coerce")
        if numeric_target.isna().any():
            raise ValueError(
                "The regression target contains non-numeric answers. Clean or convert it first."
            )
        y = numeric_target.astype(float)
        if len(y) < 10:
            raise ValueError("Regression needs at least 10 rows with target answers.")
    else:
        y = frame[target].map(str)
        counts = y.value_counts()
        if len(counts) < 2:
            raise ValueError("Classification needs at least two target classes.")
        if len(counts) > 50:
            raise ValueError("The target has too many classes for this demo (maximum 50).")
        if counts.min() < 2 or len(y) < 20:
            raise ValueError(
                "Classification needs at least 20 rows and two examples in every class."
            )
    metadata = {item["name"]: item for item in json.loads(dataset.columns_json)}
    numeric = [name for name in features if metadata[name]["type"] == "number"]
    categorical = [name for name in features if name not in numeric]
    X = frame[features].copy()
    for name in numeric:
        X[name] = pd.to_numeric(X[name].where(~X[name].map(_missing)), errors="coerce")
    for name in categorical:
        X[name] = X[name].map(lambda value: np.nan if _missing(value) else str(value))
        unique = X[name].nunique(dropna=True)
        if unique > MAX_CATEGORIES:
            raise ValueError(
                f"Feature '{name}' has {unique} categories. Exclude this identifier/free-text column."
            )
    for name in features:
        if X[name].notna().sum() == 0:
            raise ValueError(f"Feature '{name}' has no usable values. Exclude it.")
        if X[name].nunique(dropna=True) < 2:
            raise ValueError(f"Feature '{name}' is constant. Exclude it.")
    stratify = y if task == "classification" else None
    try:
        X_train, X_test, y_train, y_test = train_test_split(
            X, y, test_size=config["testSize"], random_state=seed, stratify=stratify
        )
    except ValueError as error:
        raise ValueError(f"The held-out split is not valid for this target: {error}") from error
    results = []
    artifacts = []
    for algorithm in config["algorithms"]:
        pipeline = _pipeline(task, algorithm, numeric, categorical, seed)
        pipeline.fit(X_train, y_train)
        predicted = pipeline.predict(X_test)
        common = {
            "algorithm": algorithm,
            "trainRows": len(X_train),
            "testRows": len(X_test),
            "targetRowsOmitted": omitted,
            "featureImportance": _importance(
                pipeline,
                features,
                X_test.reset_index(drop=True),
                y_test.reset_index(drop=True),
                task,
                seed,
            ),
            "importanceMethod": "Held-out permutation importance (up to 1,000 rows, 2 repeats)",
        }
        if task == "classification":
            labels = sorted({str(value) for value in y})
            majority = float(y_test.value_counts(normalize=True).max())
            metrics = {
                "accuracy": float(accuracy_score(y_test, predicted)),
                "precisionMacro": float(
                    precision_score(y_test, predicted, average="macro", zero_division=0)
                ),
                "recallMacro": float(
                    recall_score(y_test, predicted, average="macro", zero_division=0)
                ),
                "f1Macro": float(f1_score(y_test, predicted, average="macro", zero_division=0)),
                "baselineAccuracy": majority,
            }
            result = {
                **common,
                "metrics": metrics,
                "averaging": "macro",
                "classLabels": labels,
                "confusionMatrix": confusion_matrix(
                    y_test.map(str), pd.Series(predicted).map(str), labels=labels
                ).tolist(),
                "trainClassDistribution": _distribution(y_train),
                "testClassDistribution": _distribution(y_test),
            }
        else:
            rmse = float(mean_squared_error(y_test, predicted) ** 0.5)
            metrics = {
                "mae": float(mean_absolute_error(y_test, predicted)),
                "rmse": rmse,
                "r2": float(r2_score(y_test, predicted)) if len(y_test) > 1 else None,
                "baselineMae": float(
                    mean_absolute_error(y_test, np.repeat(y_train.mean(), len(y_test)))
                ),
            }
            result = {
                **common,
                "metrics": metrics,
                "actualVsPredicted": [
                    {"actual": float(actual), "predicted": float(estimate)}
                    for actual, estimate in list(zip(y_test, predicted, strict=True))[:100]
                ],
            }
        artifact_meta = {
            "task": task,
            "algorithm": algorithm,
            "target": target,
            "features": features,
            "numericFeatures": numeric,
            "categoricalFeatures": categorical,
            "versions": {
                "python": platform.python_version(),
                "scikitLearn": sklearn.__version__,
                "pandas": pd.__version__,
                "numpy": np.__version__,
                "scipy": scipy.__version__,
                "joblib": joblib.__version__,
            },
            "selected": False,
        }
        results.append(result)
        artifacts.append((algorithm, pipeline, artifact_meta))
    return results, artifacts


def execute_run(run_id: str) -> None:
    engine = make_engine(get_settings().database_url)
    SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
    written: list[Path] = []
    try:
        with SessionLocal() as db:
            run = db.get(TrainingRun, run_id)
            if run is None or run.status != "queued":
                return
            run.status = "running"
            db.commit()
            config_record = db.get(TrainingConfig, run.config_id)
            dataset = db.get(Dataset, config_record.dataset_id)
            config = json.loads(config_record.config_json)
            try:
                results, artifacts = train_models(dataset, config)
                model_dir = get_settings().model_dir
                model_dir.mkdir(parents=True, exist_ok=True)
                for index, (algorithm, pipeline, metadata) in enumerate(artifacts):
                    name = f"{uuid4().hex}.joblib"
                    temporary = model_dir / f".{uuid4().hex}.part"
                    destination = model_dir / name
                    try:
                        joblib.dump(pipeline, temporary)
                        os.replace(temporary, destination)
                    finally:
                        temporary.unlink(missing_ok=True)
                    written.append(destination)
                    metadata["artifactName"] = name
                    model = ModelMetadata(
                        owner_id=run.owner_id, run_id=run.id, metadata_json=_json(metadata)
                    )
                    db.add(model)
                    db.flush()
                    results[index]["modelId"] = model.id
                run.result_json = _json(
                    {
                        "models": results,
                        "completedAt": utcnow().isoformat(),
                        "evaluationNotice": "Metrics use the held-out test split. Repeated model selection on it does not make it an independent final assessment.",
                    }
                )
                run.status = "completed"
                db.commit()
            except Exception as error:
                db.rollback()
                run = db.get(TrainingRun, run_id)
                run.status = "failed"
                if isinstance(error, ValueError):
                    message = str(error)[:500]
                else:
                    logger.exception("Training run %s failed", run_id)
                    message = "Training failed unexpectedly. Check the selected data and try again."
                run.result_json = _json({"error": message, "failedAt": utcnow().isoformat()})
                db.commit()
                for path in written:
                    path.unlink(missing_ok=True)
    finally:
        engine.dispose()


def submit_run(run_id: str) -> None:
    executor.submit(execute_run, run_id)


def recover_interrupted_runs() -> None:
    engine = make_engine(get_settings().database_url)
    try:
        with Session(engine) as db:
            runs = db.scalars(
                select(TrainingRun).where(TrainingRun.status.in_(["queued", "running"]))
            ).all()
            for run in runs:
                run.status = "interrupted"
                run.result_json = _json(
                    {
                        "error": "Training was interrupted by an application restart. Start a new run."
                    }
                )
            db.commit()
    finally:
        engine.dispose()


def artifact_path(metadata: dict) -> Path:
    name = metadata.get("artifactName")
    if not isinstance(name, str) or not ARTIFACT_NAME.fullmatch(name):
        raise HTTPException(409, "This saved model artifact is invalid.")
    path = get_settings().model_dir / name
    if not path.is_file():
        raise HTTPException(409, "This saved model file is missing. Train the model again.")
    version = metadata.get("versions", {}).get("scikitLearn")
    if version != sklearn.__version__:
        raise HTTPException(
            409, "This model was created with a different scikit-learn version. Train it again."
        )
    return path


def load_pipeline(metadata: dict):
    try:
        return joblib.load(artifact_path(metadata))
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(409, "This model could not be loaded. Train it again.") from error


def prediction_frame(records: list[dict], metadata: dict) -> pd.DataFrame:
    features = metadata["features"]
    for row in records:
        missing = [name for name in features if name not in row]
        if missing:
            raise HTTPException(400, f"Missing required feature columns: {', '.join(missing)}.")
    frame = pd.DataFrame([{name: row[name] for name in features} for row in records], dtype=object)
    for name in metadata["numericFeatures"]:
        original = frame[name]
        frame[name] = pd.to_numeric(original.where(~original.map(_missing)), errors="coerce")
        invalid = ~original.map(_missing) & frame[name].isna()
        if invalid.any():
            raise HTTPException(400, f"Feature '{name}' needs numeric values.")
    for name in metadata["categoricalFeatures"]:
        frame[name] = frame[name].map(lambda value: np.nan if _missing(value) else str(value))
    return frame


def predict_records(pipeline, metadata: dict, records: list[dict]) -> list[dict]:
    frame = prediction_frame(records, metadata)
    values = pipeline.predict(frame)
    probabilities = pipeline.predict_proba(frame) if hasattr(pipeline, "predict_proba") else None
    classes = [str(value) for value in pipeline.classes_] if probabilities is not None else []
    output = []
    for index, value in enumerate(values):
        result = {"prediction": value.item() if hasattr(value, "item") else value}
        if probabilities is not None:
            result["probabilities"] = {
                label: float(probability)
                for label, probability in zip(classes, probabilities[index], strict=True)
            }
        output.append(result)
    return output


def model_artifact_paths(db: Session, *, dataset_ids=None, run_id=None) -> list[Path]:
    query = select(ModelMetadata)
    if run_id:
        query = query.where(ModelMetadata.run_id == run_id)
    elif dataset_ids is not None:
        query = (
            query.join(TrainingRun, ModelMetadata.run_id == TrainingRun.id)
            .join(TrainingConfig, TrainingRun.config_id == TrainingConfig.id)
            .where(TrainingConfig.dataset_id.in_(dataset_ids))
        )
    paths = []
    for model in db.scalars(query):
        name = json.loads(model.metadata_json).get("artifactName")
        if isinstance(name, str) and ARTIFACT_NAME.fullmatch(name):
            paths.append(get_settings().model_dir / name)
    return paths
