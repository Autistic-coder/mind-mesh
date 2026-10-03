"""The API parser, not the browser preview, determines dataset content."""

from datetime import date
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from backend import dataset_files


def test_csv_bom_quoted_values_identifiers_and_complete_summary():
    rows = ["id,amount,note", '001,12,"one,two"', '002,13,"line one\nline two"']
    rows += [f'{index:03d},,"later"' for index in range(3, 32)]
    summary = dataset_files.inspect_upload("research.csv", ("\ufeff" + "\n".join(rows)).encode())
    assert summary["rowCount"] == 31
    assert len(summary["preview"]) == 25
    assert summary["preview"][0] == ["001", "12", "one,two"]
    assert summary["preview"][1][2] == "line one\nline two"
    assert summary["columns"][0]["type"] == "text"
    assert summary["columns"][1]["type"] == "number"
    assert summary["columns"][1]["missing"] == 29
    assert summary["columns"][1]["uniqueCount"] == 2


def test_csv_rejects_invalid_content_and_enforces_limits(monkeypatch):
    cases = [
        (b"", "empty"),
        (b"a,a\n1,2", "unique"),
        (b"a,A\n1,2", "unique"),
        (b"a,\n1,2", "non-empty header"),
        (b"a,b\n1", "number of columns"),
        (b'a,b\n"unclosed,2', "quoting"),
        (b"\xff\xfe", "UTF-8"),
    ]
    for contents, message in cases:
        with pytest.raises(HTTPException, match=message):
            dataset_files.inspect_upload("bad.csv", contents)
    with pytest.raises(HTTPException, match="25 MB"):
        dataset_files.inspect_upload("large.csv", b"x" * (dataset_files.MAX_BYTES + 1))
    monkeypatch.setattr(dataset_files, "MAX_ROWS", 2)
    with pytest.raises(HTTPException, match="2 rows"):
        dataset_files.inspect_upload("rows.csv", b"a\n1\n2\n3")
    monkeypatch.setattr(dataset_files, "MAX_COLUMNS", 2)
    with pytest.raises(HTTPException, match="2 columns"):
        dataset_files.inspect_upload("columns.csv", b"a,b,c\n1,2,3")
    monkeypatch.setattr(dataset_files, "MAX_COLUMNS", 10_000)
    monkeypatch.setattr(dataset_files, "MAX_CELLS", 4)
    with pytest.raises(HTTPException, match="too many cells"):
        dataset_files.inspect_upload("cells.csv", b"a,b\n1,2\n3,4")


def workbook_bytes() -> bytes:
    book = Workbook()
    first = book.active
    first.title = "First"
    first.append(["value"])
    first.append([7])
    second = book.create_sheet("Selected")
    second.append(["value", "when", "flag", "identifier"])
    second.append([5, date(2025, 1, 2), True, "0012"])
    second.append([None, date(2025, 1, 3), False, "0013"])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_xlsx_sheet_selection_dates_booleans_and_text_identifiers():
    contents = workbook_bytes()
    with pytest.raises(HTTPException, match="Choose a worksheet"):
        dataset_files.inspect_upload("book.xlsx", contents)
    with pytest.raises(HTTPException, match="available worksheet"):
        dataset_files.inspect_upload("book.xlsx", contents, "Missing")
    summary = dataset_files.inspect_upload("book.xlsx", contents, "Selected")
    assert summary["rowCount"] == 2
    assert summary["preview"][0] == [5, "2025-01-02T00:00:00", "true", "0012"]
    assert [column["type"] for column in summary["columns"]] == [
        "number",
        "date",
        "boolean",
        "category",
    ]
    assert summary["columns"][0]["missing"] == 1


def test_xlsx_rejects_formulas_and_excessive_expansion(monkeypatch):
    book = Workbook()
    sheet = book.active
    sheet.append(["value"])
    sheet.append(["=1+2"])
    output = BytesIO()
    book.save(output)
    with pytest.raises(HTTPException, match="Formula cells"):
        dataset_files.inspect_upload("formula.xlsx", output.getvalue())

    source = workbook_bytes()
    expanded = BytesIO()
    with ZipFile(BytesIO(source)) as original, ZipFile(expanded, "w", ZIP_DEFLATED) as target:
        for member in original.infolist():
            target.writestr(member, original.read(member))
        target.writestr("extra.txt", "x" * 5_000)
    monkeypatch.setattr(dataset_files, "MAX_XLSX_UNCOMPRESSED", 4_000)
    with pytest.raises(HTTPException, match="too large"):
        dataset_files.inspect_upload("expanded.xlsx", expanded.getvalue(), "Selected")


def test_mixed_types_are_reported_without_changing_original_text():
    summary = dataset_files.inspect_upload("mixed.csv", b"value\n12\n0012\ntext")
    assert summary["columns"][0]["type"] == "mixed"
    assert summary["preview"] == [["12"], ["0012"], ["text"]]


def test_distinct_tracking_is_bounded_and_marked(monkeypatch):
    monkeypatch.setattr(dataset_files, "MAX_TRACKED_UNIQUE", 2)
    summary = dataset_files.inspect_upload("many.csv", b"value\none\ntwo\nthree")
    assert summary["columns"][0]["uniqueCount"] == 2
    assert summary["columns"][0]["uniqueCountCapped"] is True
    assert summary["columns"][0]["type"] == "text"
