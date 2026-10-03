"""Bounded, server-authoritative inspection of CSV and XLSX uploads."""

import csv
import io
import math
import re
import zipfile
from datetime import date, datetime
from pathlib import PurePosixPath, PureWindowsPath
from xml.etree.ElementTree import ParseError

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 200_000
MAX_COLUMNS = 10_000
MAX_PREVIEW = 25
MAX_XLSX_UNCOMPRESSED = 160 * 1024 * 1024
MAX_CELLS = 10_000_000
MAX_TRACKED_UNIQUE = 100_000
NUMBER = re.compile(r"-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?\Z")


def validate_filename(filename: str | None) -> tuple[str, str]:
    if not filename or len(filename) > 255 or not filename.isprintable():
        raise HTTPException(400, "Choose a file with a valid name.")
    if PurePosixPath(filename).name != filename or PureWindowsPath(filename).name != filename:
        raise HTTPException(400, "File names cannot contain a path.")
    dot = filename.rfind(".")
    if dot < 1 or filename[dot:].lower() not in {".csv", ".xlsx"}:
        raise HTTPException(400, "Choose a CSV or XLSX file.")
    return filename[:dot], filename[dot:].lower()


def _cell(raw) -> tuple[str | int | float, str | None]:
    # Reject formulas even when a cached result exists. Uncached formulas would
    # otherwise silently become blank cells with data_only=True.
    if hasattr(raw, "data_type") and hasattr(raw, "value"):
        if raw.data_type == "f":
            raise HTTPException(
                400, "Formula cells are not supported. Export values and upload again."
            )
        if raw.data_type == "e":
            raise HTTPException(400, "Spreadsheet error cells are not supported.")
        raw = raw.value
    if raw is None:
        return "", None
    if isinstance(raw, bool):
        return str(raw).lower(), "boolean"
    if isinstance(raw, (date, datetime)):
        return raw.isoformat(), "date"
    if isinstance(raw, (int, float)):
        if not math.isfinite(raw):
            raise HTTPException(400, "Non-finite numbers are not supported.")
        return raw, "number"
    if isinstance(raw, str):
        if not raw.strip():
            return raw, None
        # Preserve CSV text. Leading-zero identifiers and ambiguous dates stay text.
        return raw, "number" if NUMBER.fullmatch(raw) else "text"
    return str(raw), "text"


def _summarize(name: str, rows) -> dict:
    headers: list[str] | None = None
    preview: list[list[str | int | float]] = []
    missing: list[int] = []
    kinds: list[set[str]] = []
    unique: list[set[str]] = []
    examples: list[list[str]] = []
    capped: list[bool] = []
    tracked_unique = 0
    row_count = 0
    processed_cells = 0
    for raw in rows:
        processed_cells += len(raw)
        if processed_cells > MAX_CELLS:
            raise HTTPException(400, "This dataset has too many cells to inspect safely.")
        if len(raw) > MAX_COLUMNS:
            raise HTTPException(400, f"Use a dataset with {MAX_COLUMNS:,} columns or fewer.")
        values_and_kinds = [_cell(cell) for cell in raw]
        row = [value for value, _kind in values_and_kinds]
        if not any(str(cell).strip() for cell in row):
            continue
        if headers is None:
            headers = [str(cell).strip() for cell in row]
            if not headers:
                raise HTTPException(400, "Include a header row.")
            if any(not header for header in headers):
                raise HTTPException(400, "Every column needs a non-empty header.")
            if len({header.casefold() for header in headers}) != len(headers):
                raise HTTPException(400, "Column headers must be unique (ignoring case).")
            missing = [0] * len(headers)
            kinds = [set() for _ in headers]
            unique = [set() for _ in headers]
            examples = [[] for _ in headers]
            capped = [False] * len(headers)
            continue
        if len(row) != len(headers):
            raise HTTPException(400, "Each row must match the number of columns in the header.")
        row_count += 1
        if row_count > MAX_ROWS:
            raise HTTPException(400, f"Use a dataset with {MAX_ROWS:,} rows or fewer.")
        if len(preview) < MAX_PREVIEW:
            preview.append(row)
        for index, (cell, kind) in enumerate(values_and_kinds):
            if kind is None:
                missing[index] += 1
                continue
            kinds[index].add(kind)
            key = str(cell)
            if key in unique[index] or capped[index]:
                continue
            if tracked_unique >= MAX_TRACKED_UNIQUE:
                capped[index] = True
                continue
            unique[index].add(key)
            tracked_unique += 1
            if len(examples[index]) < 30:
                examples[index].append(key)
    if headers is None or row_count == 0:
        raise HTTPException(400, "Include a header row and at least one data row.")
    columns = []
    for index, header in enumerate(headers):
        observed = kinds[index]
        if len(observed) > 1:
            inferred = "mixed"
        elif observed == {"number"}:
            inferred = "number"
        elif observed == {"date"}:
            inferred = "date"
        elif observed == {"boolean"}:
            inferred = "boolean"
        elif not capped[index] and len(unique[index]) <= 30:
            inferred = "category"
        else:
            inferred = "text"
        columns.append(
            {
                "name": header,
                "type": inferred,
                "missing": missing[index],
                "uniqueCount": len(unique[index]),
                "uniqueCountCapped": capped[index],
                "values": examples[index],
            }
        )
    return {"name": name, "rowCount": row_count, "columns": columns, "preview": preview}


def inspect_upload(filename: str, data: bytes, sheet_name: str | None = None) -> dict:
    stem, extension = validate_filename(filename)
    if not data:
        raise HTTPException(400, "This file is empty.")
    if len(data) > MAX_BYTES:
        raise HTTPException(400, "Choose a file no larger than 25 MB.")
    if extension == ".csv":
        if sheet_name:
            raise HTTPException(400, "A worksheet can only be selected for XLSX files.")
        try:
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise HTTPException(400, "This is not a valid text CSV file.")
            return _summarize(stem, csv.reader(io.StringIO(text, newline=""), strict=True))
        except UnicodeDecodeError as error:
            raise HTTPException(400, "This CSV file must use UTF-8 encoding.") from error
        except csv.Error as error:
            raise HTTPException(
                400, "This CSV file has invalid quoting or a field that is too large."
            ) from error

    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise HTTPException(400, "This is not a valid XLSX workbook.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if (
                len(members) > 2_000
                or sum(member.file_size for member in members) > MAX_XLSX_UNCOMPRESSED
            ):
                raise HTTPException(400, "This workbook is too large to inspect safely.")
            expanded = 0
            for member in members:
                with archive.open(member) as source:
                    while chunk := source.read(1024 * 1024):
                        expanded += len(chunk)
                        if expanded > MAX_XLSX_UNCOMPRESSED:
                            raise HTTPException(
                                400, "This workbook is too large to inspect safely."
                            )
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            if sheet_name and sheet_name not in workbook.sheetnames:
                raise HTTPException(400, "Choose an available worksheet.")
            if not sheet_name and len(workbook.sheetnames) > 1:
                raise HTTPException(400, "Choose a worksheet from this workbook.")
            chosen = sheet_name or workbook.sheetnames[0]
            sheet = workbook[chosen]
            if (sheet.max_column and sheet.max_column > MAX_COLUMNS) or (
                sheet.max_row and sheet.max_row > MAX_ROWS + 1
            ):
                raise HTTPException(
                    400,
                    f"Use a worksheet with {MAX_ROWS:,} rows and {MAX_COLUMNS:,} columns or fewer.",
                )
            if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > MAX_CELLS:
                raise HTTPException(400, "This worksheet has too many cells to inspect safely.")
            label = f"{stem} · {chosen}" if len(workbook.sheetnames) > 1 else stem
            return _summarize(label, sheet.iter_rows(values_only=False))
        finally:
            workbook.close()
    except (
        ValueError,
        TypeError,
        KeyError,
        zipfile.BadZipFile,
        OSError,
        IndexError,
        InvalidFileException,
        ParseError,
    ) as error:
        raise HTTPException(400, "This XLSX workbook could not be read.") from error
